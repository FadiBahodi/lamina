"""Mechanical source decomposition with stable identity and honest locators."""

from __future__ import annotations

import hashlib
import re
import subprocess
import tempfile
from pathlib import Path

from .store import Workspace

SUPPORTED = {".md", ".txt", ".pdf", ".pptx"}


def _paragraphs(text: str):
    """Keep complete paragraphs, including long multilingual paragraphs."""
    start = 0
    for match in re.finditer(r"\n[ \t]*\n", text):
        value = text[start : match.start()].strip()
        if value:
            yield start, match.start(), value, "paragraph"
        start = match.end()
    if text[start:].strip():
        yield start, len(text), text[start:].strip(), "paragraph"


def _text_sections(text: str, title: str):
    """Read heading ancestry with Markdown's parser, including fenced code."""
    from markdown_it import MarkdownIt

    lines = text.splitlines()
    tokens = MarkdownIt("commonmark").parse(text)
    heading, start, ancestry = title, 0, []
    for index, token in enumerate(tokens):
        if token.type != "heading_open" or token.level != 0:
            continue
        lo, hi = token.map
        body = "\n".join(lines[start:lo])
        if body.strip():
            yield heading, body, start + 1, [name for _, name in ancestry] or [title]
        heading = tokens[index + 1].content
        level = int(token.tag[1:])
        ancestry = [(depth, name) for depth, name in ancestry if depth < level]
        ancestry.append((level, heading))
        start = hi
    body = "\n".join(lines[start:])
    if body.strip():
        yield heading, body, start + 1, [name for _, name in ancestry] or [title]


def _structured_chunks(text):
    """Return complete parser blocks; capacity is decided when building requests.

    A list, table, block quote, code fence, or paragraph stays intact regardless
    of character count. Parser gaps (including reference definitions) remain in
    the neighboring raw source span, so ingestion never drops their text.
    """
    from markdown_it import MarkdownIt

    lines = text.splitlines(keepends=True)
    offsets = [0]
    for line in lines:
        offsets.append(offsets[-1] + len(line))
    tokens = MarkdownIt("commonmark").enable("table").parse(text)
    spans = [
        (offsets[t.map[0]], offsets[t.map[1]], t.type.removesuffix("_open"))
        for t in tokens
        if t.level == 0 and t.map and t.nesting != -1
    ]
    previous = 0
    for index, (_, hi, kind) in enumerate(spans):
        lo, previous = previous, hi
        if index == len(spans) - 1:
            hi = len(text)
        value = text[lo:hi].strip()
        if value:
            yield lo, hi, value, kind
    if not spans and text.strip():
        yield 0, len(text), text.strip(), "definitions"


def _reference_context(text, references, by_destination):
    """Keep Markdown's resolved links available across section boundaries."""
    if not references:
        return []
    from markdown_it import MarkdownIt

    tokens = MarkdownIt("commonmark").parse(text, {"references": references})
    destinations = {
        (child.attrGet("href"), child.attrGet("title") or "")
        for token in tokens
        for child in (token.children or [])
        if child.type == "link_open"
    }
    return [
        entry
        for destination in sorted(destinations)
        for entry in by_destination.get(destination, [])
    ]


def _pptx_chunks(path):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ValueError(
            'PowerPoint support requires: pip install "lamina-engine[slides]"'
        ) from exc
    from zipfile import ZipFile

    with ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100_000_000:
            raise ValueError("PPTX expanded contents exceed 100 MB; split the deck")
    for number, slide in enumerate(Presentation(path).slides, 1):
        heading = slide.shapes.title.text if slide.shapes.title else f"Slide {number}"
        found = False

        def walk(shapes):
            for shape in sorted(shapes, key=lambda s: (s.top, s.left)):
                if hasattr(shape, "shapes"):
                    yield from walk(shape.shapes)
                else:
                    yield shape

        for shape in walk(slide.shapes):
            if shape.has_table:
                # Row/column boundaries remain explicit; cells are not flattened.
                rows = [
                    [
                        cell.text.replace("\n", " / ").replace("|", "\\|")
                        for cell in row.cells
                    ]
                    for row in shape.table.rows
                ]
                value = "\n".join(" | ".join(row) for row in rows)
                kind = "table"
            elif shape.has_text_frame:
                value = "\n".join(
                    "  " * p.level + p.text for p in shape.text_frame.paragraphs
                )
                kind = "text"
            else:
                continue
            if value.strip():
                found = True
                yield heading, value.strip(), f"slide {number}, {kind} {shape.shape_id}", {
                    "kind": kind,
                    "slide": number,
                    "shape": shape.shape_id,
                    "structural_group": f"slide:{number}",
                    "heading_path": [heading],
                }
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text.strip():
                found = True
                yield heading, frame.text.strip(), f"slide {number}, speaker notes", {
                    "kind": "speaker_notes",
                    "slide": number,
                    "structural_group": f"slide:{number}",
                    "heading_path": [heading],
                }
        if not found:
            raise ValueError(
                f"{path.name}, slide {number}: no extractable text; inspect visuals or convert and OCR first"
            )


def read_source(
    path: Path, role: str = "teaching", *, ocr: bool = False
) -> tuple[dict, list[dict]]:
    if role not in {"teaching", "assessment"}:
        raise ValueError("role must be teaching or assessment")
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError(
            f"Unsupported input: {path.name}; use Markdown, text, PDF, or PPTX (convert legacy .ppt first)"
        )
    raw = path.read_bytes()
    checksum = hashlib.sha256(raw).hexdigest()
    source_id = "src_" + checksum[:20]
    title = path.stem.replace("_", " ").replace("-", " ").strip()
    chunks = []
    if path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError(
                'PDF support requires: pip install "lamina-engine[pdf]"'
            ) from exc
        # Explicit local preprocessing. Native text pages remain untouched.
        # The source checksum always identifies the original input file.
        if ocr:
            with tempfile.TemporaryDirectory(prefix="lamina-ocr-") as temporary:
                output = Path(temporary) / "searchable.pdf"
                try:
                    subprocess.run(
                        [
                            "ocrmypdf",
                            "--skip-text",
                            "--",
                            str(path.resolve()),
                            str(output),
                        ],
                        check=True,
                        capture_output=True,
                        timeout=1800,
                    )
                except FileNotFoundError as exc:
                    raise ValueError(
                        "Install OCRmyPDF and its Tesseract dependency to use --ocr"
                    ) from exc
                except (
                    subprocess.CalledProcessError,
                    subprocess.TimeoutExpired,
                ) as exc:
                    raise ValueError(
                        f"OCR failed for {path.name}; original input was not changed"
                    ) from exc
                from io import BytesIO

                reader = PdfReader(BytesIO(output.read_bytes()))
        else:
            reader = PdfReader(path)
        for page_number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise ValueError(
                    f"{path.name}, page {page_number}: no extractable text; OCR first (image-only pages are not silently skipped)"
                )
            chunks.append(
                (
                    title,
                    text.strip(),
                    f"page {page_number}",
                    {
                        "kind": "page",
                        "page": page_number,
                        "structural_group": f"page:{page_number}",
                        "heading_path": [title],
                    },
                )
            )
    elif path.suffix.lower() == ".pptx":
        chunks.extend(_pptx_chunks(path))
    else:
        text = raw.decode("utf-8-sig")
        reference_env = {}
        reference_index = {}
        if path.suffix.lower() == ".md":
            from markdown_it import MarkdownIt

            tokens = MarkdownIt("commonmark").parse(text, reference_env)
            title = next(
                (
                    tokens[index + 1].content
                    for index, token in enumerate(tokens)
                    if token.type == "heading_open"
                    and token.tag == "h1"
                    and token.level == 0
                ),
                title,
            )
            for label, entry in reference_env.get("references", {}).items():
                reference_index.setdefault(
                    (entry["href"], entry.get("title", "")), []
                ).append(
                    {
                        "label": label,
                        "href": entry["href"],
                        "title": entry.get("title", ""),
                    }
                )
        for section_index, (heading, body, first_line, heading_path) in enumerate(
            _text_sections(text, title)
        ):
            for start, end, value, kind in (
                _structured_chunks(body)
                if path.suffix.lower() == ".md"
                else _paragraphs(body)
            ):
                lo = first_line + body[:start].count("\n")
                hi = first_line + body[:end].count("\n")
                chunks.append(
                    (
                        heading,
                        value,
                        f"lines {lo}–{hi}",
                        {
                            "kind": kind,
                            "heading_path": heading_path,
                            "section_id": "section_"
                            + hashlib.sha256(
                                "\n".join(heading_path).encode()
                            ).hexdigest()[:16],
                            "section_index": section_index,
                            **(
                                {"reference_context": references}
                                if (
                                    references := _reference_context(
                                        value,
                                        reference_env.get("references", {}),
                                        reference_index,
                                    )
                                )
                                else {}
                            ),
                        },
                    )
                )
    if not chunks:
        raise ValueError(f"{path.name} has no readable content")
    source = {
        "id": source_id,
        "title": title,
        "filename": path.name,
        "sha256": checksum,
        "role": role,
        "parser": "structured-text-3",
        "ocr": bool(ocr and path.suffix.lower() == ".pdf"),
        "limits": "Text extraction does not interpret diagrams, charts, or image meaning.",
    }
    units = []
    for ordinal, (heading, text, locator, structure) in enumerate(chunks):
        unit_id = (
            "unit_"
            + hashlib.sha256(f"{source_id}:{ordinal}:{text}".encode()).hexdigest()[:20]
        )
        units.append(
            {
                "id": unit_id,
                "source_id": source_id,
                "heading": heading,
                "text": text,
                "locator": locator,
                "ordinal": ordinal,
                "role": role,
                "content_id": hashlib.sha256(
                    ("\n".join(structure["heading_path"]) + "\n" + text).encode()
                ).hexdigest(),
                **structure,
            }
        )
    return source, units


def ingest_paths(
    paths: list[Path],
    workspace: Workspace,
    role: str = "teaching",
    *,
    ocr: bool = False,
) -> dict:
    discovered = []
    for input_path in paths:
        path = Path(input_path)
        if not path.exists():
            raise ValueError(f"Input does not exist: {path}")
        if path.is_dir():
            discovered.extend(
                p
                for p in sorted(path.rglob("*"))
                if p.is_file() and p.suffix.lower() in SUPPORTED
            )
        else:
            discovered.append(path)
    if not discovered:
        raise ValueError("No supported files found")
    # Parse the entire batch before storing anything, so extraction failures
    # cannot leave a deceptively complete partial import.
    parsed = [read_source(p, role, ocr=ocr) for p in discovered]
    existing = {s["id"]: s for s in workspace.sources()}
    for source, _ in parsed:
        if source["id"] in existing and existing[source["id"]]["role"] != role:
            raise ValueError(
                "Identical content already exists under another source role"
            )
    new_sources = 0
    for source, units in parsed:
        if source["id"] not in existing:
            new_sources += 1
            existing[source["id"]] = source
    workspace.import_sources(parsed)
    return {
        "files_read": len(parsed),
        "new_sources": new_sources,
        "duplicates": len(parsed) - new_sources,
        **workspace.stats(),
    }
