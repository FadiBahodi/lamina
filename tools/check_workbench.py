"""Connected browser smoke test with deterministic model responses.

Run: PYTHONPATH=src:tests python tools/check_workbench.py
Requires playwright with Chromium, python-pptx, and the normal test dependencies.
"""

import json
import tempfile
import threading
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation

from lamina.studio_server import StudioServer
from test_source_workflows import SourceWriter


def main():
    with tempfile.TemporaryDirectory(prefix="lamina-browser-") as folder:
        root = Path(folder)
        source = root / "wind.md"
        source.write_text("# Wind\n\nWind transfers energy to the blades.")
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Shaft"
        slide.placeholders[1].text = "The shaft transfers rotation to a generator."
        slides = root / "wind.pptx"
        deck.save(slides)
        server = StudioServer(root / "workspace", port=0)
        server.provider = SourceWriter()
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True, args=["--no-sandbox"])
                page = browser.new_page(viewport={"width": 1280, "height": 960})
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                origin = f"http://127.0.0.1:{server.server_port}"
                page.goto(origin)
                page.locator('[data-view="workflows"]').first.click()
                page.get_by_label("Project goal").fill(
                    "Explain how a wind turbine generates electricity."
                )
                page.locator(
                    '#production-workbench input[type="file"]'
                ).set_input_files([str(source), str(slides)])
                page.wait_for_function(
                    "document.querySelector('.pb-upload-status').textContent.includes('Stored 2 sources')"
                )
                page.get_by_role("button", name="Start project", exact=True).click()
                page.get_by_role("link", name="Download text", exact=True).wait_for(
                    timeout=30000
                )
                href = page.get_by_role(
                    "link", name="Download text", exact=True
                ).get_attribute("href")
                response = page.request.get(href)
                assert response.ok and "shaft" in response.text().lower()
                projects = page.request.get(origin + "/api/projects").json()
                # Both progress and full results are checked against the actual server.
                rid = next(iter(server.runs))
                progress = page.request.get(origin + f"/api/progress/{rid}").json()
                assert "plan" not in progress and "receipt" not in progress
                result = page.request.get(origin + f"/api/runs/{rid}").json()
                assert result["plan"]["options"]["workflow"] == "direct"
                assert len(result["receipt"]["metrics"]["requests"]) == 2
                page.get_by_label("Requested section change").fill(
                    "Clarify the energy transfer."
                )
                page.get_by_role("button", name="Revise section", exact=True).click()
                page.get_by_role("link", name="Download text", exact=True).wait_for(
                    timeout=30000
                )
                page.wait_for_function(
                    "document.querySelector('.pb-output').textContent.includes('Clarify the energy transfer.')"
                )
                assert not errors, errors
                print(
                    json.dumps(
                        {
                            "kind": "deterministic_browser_smoke",
                            "checks": [
                                "Markdown and PowerPoint upload",
                                "automatic direct writing",
                                "download source-backed text",
                                "compact progress",
                                "targeted revision",
                                "no browser script errors",
                            ],
                        }
                    )
                )
                browser.close()
        finally:
            server.shutdown()
            thread.join(timeout=5)
            server.server_close()


if __name__ == "__main__":
    main()
