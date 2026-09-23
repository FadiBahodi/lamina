"""Reproducible machinery evaluation on five original PDFs and three PPTX files.

The provider below is a fixture with access to the corpus oracle. It cannot
measure model intelligence, extraction recall on unknown text, cost, or live
provider latency. Small seeded lognormal sleeps exercise scheduling variance.
Run: python benchmarks/workflow_matrix.py --output benchmarks/results/workflow-matrix.json
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import random
import re
import sys
import tempfile
import threading
import time

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from lamina.ingest import ingest_paths
from lamina.production import build_production
from lamina.production_example import fixture_request_budget
from lamina.store import Workspace, canonical

# These are invented equipment specifications, deliberately carrying a
# condition, negation, historical conflict, duplicate, table and notes-only fact.
# They must never be presented as engineering guidance for real equipment.
FACTS = {
    "mechanics": [
        "[F01] Rotor M uses three blades to transfer rotation to its generator.",
        "[F02] Rotor M stops above 20 m/s unless supervised test mode is active.",
        "[F03] Rotor M does not produce electricity while its brake is locked.",
    ],
    "control": [
        "[F04] Controller C starts after the sensor agrees with its backup for 30 seconds.",
        "[F05] Controller C enters standby when the two wind sensors disagree.",
        "[F06] A standby alarm alone does not prove that a wind sensor is broken.",
    ],
    "history": [
        "[F07] The 2010 Rotor M handbook states a shutdown threshold of 18 m/s.",
        "[F08] The 2025 Rotor M bulletin replaces the 2010 threshold with 20 m/s.",
        "[F09] The 2025 bulletin applies only to Rotor M controllers with firmware C.",
    ],
    "inspection": [
        "[F10] Inspect the blade root after a recorded overload before restarting.",
        "[F11] A missing overload record does not establish that no overload occurred.",
        "[F12] During inspection the isolation switch stays open and tagged.",
    ],
    "measurement": [
        "[F13] Channel A measures wind speed in m/s; Channel B measures rotor speed in rpm.",
        "[F14] A negative Channel A value is invalid; preserve the raw value in the log.",
        "[F15] The table gives fixture readings, not operating limits: A=12 m/s, B=120 rpm.",
    ],
}
QUALIFIED = "Rotor M stops above 20 m/s unless supervised test mode is active."
UNQUALIFIED = "Rotor M always stops above 20 m/s."
SECRET = "HELD_OUT_FIXTURE_KEY_941"


def normalized(value):
    return " ".join(value.split())


def make_corpus(folder: Path, *, revised=False) -> dict:
    """Create originals through the same libraries users' inputs exercise."""
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import simpleSplit
    from pptx import Presentation
    from pptx.util import Inches

    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for index, (topic, facts) in enumerate(FACTS.items(), 1):
        path = folder / f"{index:02d}_{topic}.pdf"
        canvas = Canvas(str(path), pagesize=(612, 792), invariant=1)
        canvas.setTitle(f"Original fixture: {topic}")
        for number, fact in enumerate(facts, 1):
            if revised and topic == "measurement" and number == 2:
                fact += " Record the calibration version as well."
            canvas.setFont("Helvetica-Bold", 14)
            canvas.drawString(
                48, 738, f"[{topic}] Original windmill fixture, page {number}"
            )
            canvas.setFont("Helvetica", 11)
            for line, text in enumerate(simpleSplit(fact, "Helvetica", 11, 505)):
                canvas.drawString(48, 700 - line * 16, text)
            canvas.showPage()
        canvas.save()
        paths.append(path)
    pairs = [
        ("mechanics", "control"),
        ("history", "inspection"),
        ("measurement", "mechanics"),
    ]
    for index, pair in enumerate(pairs, 1):
        path = folder / f"deck_{index}.pptx"
        deck = Presentation()
        for topic in pair:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = f"[{topic}] Windmill workshop"
            slide.placeholders[1].text = FACTS[topic][0]
            slide.notes_slide.notes_text_frame.text = FACTS[topic][1]
            if topic == "measurement":
                table = slide.shapes.add_table(
                    3, 3, Inches(0.6), Inches(4), Inches(8), Inches(1.2)
                ).table
                for row, values in enumerate(
                    (
                        ("Channel", "Unit", "Reading"),
                        ("A", "m/s", "12"),
                        ("B", "rpm", "120"),
                    )
                ):
                    for column, text in enumerate(values):
                        table.cell(row, column).text = text
        deck.save(path)
        paths.append(path)
    # A separate held-out file verifies the source-role boundary.
    key = folder / "held_out.md"
    key.write_text("# Private examiner key\n\n" + SECRET, encoding="utf-8")
    return {"sources": paths, "held_out": key}


def topic_of(row):
    value = (
        row.get("heading", "") + " " + row.get("title", "") + " " + row.get("text", "")
    )
    for topic in FACTS:
        if f"[{topic}]" in value:
            return topic
    return "mechanics"


def model_of(row):
    match = re.search(
        r"\[M\d+\]",
        " ".join(
            str(row.get(key, "")) for key in ("heading", "title", "text", "explanation")
        ),
    )
    return match.group()[1:-1] if match else None


def assignments(units):
    groups = {}
    for unit in units:
        groups.setdefault(topic_of(unit), []).append(unit["id"])
    result = {
        "title": "Original windmill fixture",
        "summary": "Preserve fixture specifications and qualifications.",
        "sections": [
            {
                "id": topic,
                "title": topic.title(),
                "purpose": "Preserve the assigned fixture facts.",
                "unit_ids": ids,
                "context_section_ids": [],
                "representation": {
                    "kind": "reference",
                    "rationale": "Fixture accounting.",
                    "requirements": [],
                },
            }
            for topic in FACTS
            if (ids := groups.get(topic))
        ],
        "omitted": [],
        "shared_context": [],
    }
    for section in result["sections"]:
        if section["id"] == "history" and "mechanics" in groups:
            section["context_section_ids"] = ["mechanics"]
    return result


class ScriptedProvider:
    """Original source echo plus an oracle-controlled qualification defect."""

    def __init__(self, *, fault=None, delays=True):
        self.fault, self.delays = fault, delays
        self.identity = "workflow-matrix-scripted-1:" + str(fault)
        self.requests = []
        self.lock = threading.Lock()
        self.read_fault_seen = False

    def budget_for(self, stage):
        # This oracle is deterministic. Its finite matrix allows up to 512
        # items per request; this is not a deployable model-quality profile.
        return fixture_request_budget(stage, max_items=512)

    def call(self, stage, payload):
        start = time.monotonic()
        data = payload["input"]
        request_text = canonical(payload)
        if SECRET in request_text:
            raise AssertionError("Held-out source leaked into a provider request")
        with self.lock:
            record = {"stage": stage, "payload": payload, "start": start}
            self.requests.append(record)
        seed = int(hashlib.sha256(request_text.encode()).hexdigest()[:16], 16)
        if self.delays:
            # The distribution is explicitly synthetic and keyed per request,
            # so worker completion order cannot alter the random draws.
            median = 0.008 if stage == "production_write" else 0.004
            time.sleep(random.Random(seed).lognormvariate(0, 0.7) * median)
        try:
            response = self._response(stage, data)
            record["response"] = response
            return response
        finally:
            record["end"] = time.monotonic()

    def _response(self, stage, data):
        if stage == "production_read":
            with self.lock:
                fail = self.fault == "persistent-invalid-quote" or (
                    self.fault == "one-invalid-quote" and not self.read_fault_seen
                )
                self.read_fault_seen = True
            response = {
                "ideas": [
                    {
                        "title": f"[{topic_of(unit)}] "
                        + (f"[{model_of(unit)}] " if model_of(unit) else "")
                        + "Fixture source statement",
                        "explanation": unit["text"],
                        "unit_ids": [unit["id"]],
                        **(
                            {
                                "evidence": [
                                    {"unit_id": unit["id"], "quote": "INVENTED_SPAN"}
                                ]
                            }
                            if fail
                            else {
                                "evidence_refs": [
                                    {
                                        "span_id": unit["spans"][0]["id"],
                                        "end_span_id": unit["spans"][-1]["id"],
                                    }
                                ]
                            }
                        ),
                    }
                    for unit in data["core"]
                ]
            }
            if "repair" in data:
                return {
                    "replacements": [
                        {"index": row["index"], "idea": response["ideas"][row["index"]]}
                        for row in data["repair"]["invalid_ideas"]
                    ]
                }
            return response
        if stage == "production_route":
            groups = {}
            for idea in data["ideas"]:
                key = model_of(idea) if data.get("parent_section") else topic_of(idea)
                groups.setdefault(key or topic_of(idea), []).append(idea["id"])
            result = {
                "title": "Original windmill fixture",
                "summary": "Fixture accounting.",
                "sections": [
                    {
                        "id": topic,
                        "title": topic.title(),
                        "purpose": "Preserve assigned fixture facts.",
                        "idea_ids": ids,
                        "context_section_ids": [],
                        "representation": {
                            "kind": "reference",
                            "rationale": "Fixture accounting.",
                            "requirements": [],
                        },
                    }
                    for topic, ids in sorted(
                        groups.items(),
                        key=lambda pair: (
                            (
                                list(FACTS).index(pair[0])
                                if pair[0] in FACTS
                                else len(FACTS)
                            ),
                            pair[0],
                        ),
                    )
                ],
                "omitted": [],
                "shared_context": [],
            }
            for section in result["sections"]:
                if section["id"] == "history" and "mechanics" in groups:
                    section["context_section_ids"] = ["mechanics"]
            return result
        if stage == "production_group":
            groups = {}
            for card in data["cards"]:
                key = model_of(card) if data.get("parent_section") else topic_of(card)
                groups.setdefault(key or topic_of(card), []).append(card["id"])
            return {
                "groups": [
                    {
                        "id": topic,
                        "title": f"[{topic}] Fixture source family",
                        "explanation": f"Fixture statements concerning {topic}.",
                        "member_ids": ids,
                    }
                    for topic, ids in groups.items()
                ]
            }
        if stage == "production_assign":
            sections = data["outline"]["sections"]
            groups = {}
            for idea in data["ideas"]:
                topic = (
                    model_of(idea)
                    if any(s["id"] == model_of(idea) for s in sections)
                    else topic_of(idea)
                )
                section = next((s for s in sections if s["id"] == topic), sections[0])
                groups.setdefault(section["id"], []).append(idea["id"])
            return {
                "assignments": [
                    {"section_id": sid, "idea_ids": ids} for sid, ids in groups.items()
                ],
                "omitted": [],
                "shared_context": [],
            }
        if stage in {"production_write", "production_repair"}:
            units = data["assigned_units"]
            body = "\n\n".join(unit["text"] for unit in units)
            if stage == "production_write":
                body = body.replace(QUALIFIED, UNQUALIFIED)
            # Exact body/source mappings remain valid for the defective claim;
            # the separate review has to notice its dropped qualification.
            claims, evidence = [], []
            for i, unit in enumerate(units):
                text = (
                    unit["text"].replace(QUALIFIED, UNQUALIFIED)
                    if stage == "production_write"
                    else unit["text"]
                )
                source = {"unit_id": unit["id"], "quote": unit["text"]}
                claims.append(
                    {
                        "id": f"c{i}",
                        "text": text,
                        "body_field": "body",
                        "evidence": [source],
                    }
                )
                evidence.append(source)
            result = {"body": body, "evidence": evidence, "claims": claims}
            section = data["section"]
            if "unit_ids" in section:
                result.update(used_unit_ids=section["unit_ids"], omitted_units=[])
            else:
                result["used_idea_ids"] = section["idea_ids"]
            if data["format"] == "assessment":
                result.pop("body")
                result["candidate_body"] = (
                    f"What conditions apply to {section['title']} in the fixture?"
                )
                result["marking_body"] = body
                for claim in claims:
                    claim["body_field"] = "marking_body"
            return result
        if stage == "production_review":
            draft = data["draft"]
            if UNQUALIFIED in draft.get("body", draft.get("marking_body", "")):
                unit = next(
                    unit for unit in data["assigned_units"] if QUALIFIED in unit["text"]
                )
                return {
                    "findings": [
                        {
                            "issue": "The supervised test-mode exception disappeared.",
                            "repair_instruction": "Restore the exact supervised test-mode exception.",
                            "evidence": [
                                {"unit_id": unit["id"], "quote": unit["text"]}
                            ],
                        }
                    ]
                }
            return {"findings": []}
        if stage == "production_consistency":
            # The fixture explicitly preserves both dates and their different
            # thresholds; declaring agreement would be an oracle failure.
            return {"findings": []}
        if stage == "assessment_blind_solve":
            assert set(data) == {"current_candidate_body", "prior_candidate_bodies"}
            assert "[F" not in canonical(
                data
            ), "The blind solve received the source or marking key"
            return {
                "answer": "The fixture condition and its stated exception apply.",
                "uncertainties": [],
            }
        if stage == "assessment_judge":
            return {"findings": []}
        raise ValueError(f"Fixture provider does not implement {stage}")


def measure(
    workspace,
    provider,
    source_ids,
    workflow,
    *,
    brief="Preserve all fixture specifications, conflicts, and conditions.",
    required_facts=None,
    extra_options=None,
):
    events = []
    start = time.monotonic()
    offset = len(provider.requests)

    def progress(event):
        events.append({**event, "elapsed_ms": (time.monotonic() - start) * 1000})

    options = {
        "workflow": workflow,
        "workers": 4,
        "max_input_bytes": 120_000,
        # The matrix explicitly evaluates reusable-reading cache mechanics.
        "reading": "reusable",
    }
    options.update(extra_options or {})
    if workflow == "assigned":
        options["assignments"] = assignments(workspace.units("teaching", source_ids))
    receipt, error = None, None
    try:
        receipt = build_production(
            workspace, provider, brief, source_ids, options, progress
        )
    except Exception as exc:
        error = {"type": type(exc).__name__, "message": str(exc)}
    elapsed = (time.monotonic() - start) * 1000
    requests = provider.requests[offset:]
    counts = dict(Counter(row["stage"] for row in requests))
    sizes = [len(canonical(row["payload"]).encode()) for row in requests]
    active, peak = 0, 0
    for _, delta in sorted(
        [(row["start"], 1) for row in requests] + [(row["end"], -1) for row in requests]
    ):
        active += delta
        peak = max(peak, active)
    stage_times = {}
    for stage in counts:
        selected = [row for row in requests if row["stage"] == stage]
        stage_times[stage] = {
            "first_provider_start_ms": round(
                (min(row["start"] for row in selected) - start) * 1000, 3
            ),
            "last_provider_completion_ms": round(
                (max(row["end"] for row in selected) - start) * 1000, 3
            ),
            "summed_provider_occupancy_ms": round(
                sum((row["end"] - row["start"]) * 1000 for row in selected), 3
            ),
        }
    try:
        import tiktoken

        encoding = tiktoken.get_encoding("cl100k_base")
        tokens = sum(
            len(encoding.encode(canonical(row["payload"]))) for row in requests
        )
        output_tokens = sum(
            len(encoding.encode(canonical(row["response"])))
            for row in requests
            if "response" in row
        )
    except ImportError:
        tokens = output_tokens = None
    completed_drafts = [
        e["elapsed_ms"]
        for e in events
        if e["stage"] == "production_write" and e["status"] == "completed"
    ]
    completed_reviews = [
        e["elapsed_ms"]
        for e in events
        if e["stage"] == "production_review" and e["status"] == "completed"
    ]
    terminal_review = {}
    for event in events:
        if event["stage"] == "production_review" and event["status"] == "completed":
            terminal_review[event["item"]] = event["elapsed_ms"]
    source_words = sum(
        len(u["text"].split()) for u in workspace.units("teaching", source_ids)
    )
    read_words = sum(
        len(unit["text"].split())
        for row in requests
        if row["stage"] == "production_read"
        for group in ("core", "before", "after")
        for unit in row["payload"]["input"].get(group, [])
    )
    result = {
        "status": receipt["status"] if receipt else "failed",
        "reading_policy": options["reading"],
        "fixture_workload": {
            "basis": "configured",
            "max_items_per_stage": provider.budget_for(
                "production_read"
            ).workload.max_items,
            "purpose": "Finite scripted oracle; no model-quality evidence",
        },
        "error": error,
        "calls": len(requests),
        "calls_by_stage": counts,
        "stage_times": stage_times,
        "peak_simultaneous_provider_calls": peak,
        "request_bytes": sum(sizes),
        "largest_request_bytes": max(sizes, default=0),
        "reference_input_tokens_cl100k_base": tokens,
        "reference_output_tokens_cl100k_base": output_tokens,
        "response_bytes": sum(
            len(canonical(row["response"]).encode())
            for row in requests
            if "response" in row
        ),
        "provider_reported_usage": None,
        "synthetic_end_to_end_ms": round(elapsed, 3),
        "first_draft_completed_ms": (
            round(min(completed_drafts), 3) if completed_drafts else None
        ),
        "first_review_completed_ms": (
            round(min(completed_reviews), 3) if completed_reviews else None
        ),
        "first_terminal_section_review_ms": (
            round(min(terminal_review.values()), 3)
            if receipt and terminal_review
            else None
        ),
        "summed_provider_occupancy_ms": round(
            sum((r["end"] - r["start"]) * 1000 for r in requests), 3
        ),
        "reader_source_word_amplification": (
            round(read_words / source_words, 3) if source_words else None
        ),
    }
    if receipt:
        output = normalized(receipt.get("examiner_markdown") or receipt["markdown"])
        expected = (
            required_facts
            if required_facts is not None
            else [fact for facts in FACTS.values() for fact in facts]
        )
        missing = [fact for fact in expected if normalized(fact) not in output]
        result["fixture_oracle"] = {
            "required_fact_count": len(expected),
            "missing_fact_ids": [
                re.search(r"\[[A-Z]\d+\]", fact).group() for fact in missing
            ],
            "qualification_present": (
                QUALIFIED in output if required_facts is None else None
            ),
            "false_unqualified_statement_present": UNQUALIFIED in output,
            "both_historical_thresholds_present": "18 m/s" in output
            and "20 m/s" in output,
            "held_out_key_present": SECRET in output,
        }
        result["coverage"] = receipt.get("coverage")
        result["output_sha256"] = hashlib.sha256(
            (receipt["markdown"] + (receipt.get("examiner_markdown") or "")).encode()
        ).hexdigest()
        result["plan_phase_ms"] = receipt["plan"]["metrics"]["wall_ms"]
        result["production_phase_ms"] = receipt["metrics"]["wall_ms"]
        result["document_checks"] = receipt.get("document_checks")
        result["assessment_checks"] = (
            {
                "status": receipt["assessment_checks"]["status"],
                "candidate_contains_source_fact_markers": "[F"
                in receipt["candidate_markdown"],
            }
            if receipt.get("assessment_checks")
            else None
        )
    return result


def run_matrix(folder):
    corpus = make_corpus(folder / "corpus")
    cases = {}
    parse_metrics = None
    for workflow in ("direct", "planned", "assigned"):
        workspace = Workspace(folder / workflow)
        started = time.monotonic()
        ingest_paths(corpus["sources"], workspace)
        parse_elapsed = (time.monotonic() - started) * 1000
        ingest_paths([corpus["held_out"]], workspace, role="assessment")
        source_ids = [
            row["id"] for row in workspace.sources() if row["role"] == "teaching"
        ]
        parse_metrics = {
            "pdf_files": 5,
            "pptx_files": 3,
            "teaching_units": len(workspace.units("teaching")),
            "last_import_ms": round(parse_elapsed, 3),
        }
        provider = ScriptedProvider()
        cases[f"{workflow}/cold"] = measure(workspace, provider, source_ids, workflow)
        cases[f"{workflow}/warm"] = measure(workspace, provider, source_ids, workflow)
        cases[f"{workflow}/new-brief"] = measure(
            workspace,
            provider,
            source_ids,
            workflow,
            brief="Create a concise reference preserving every fixture fact and qualification.",
        )
        revised = make_corpus(folder / "revised", revised=True)
        changed = next(p for p in revised["sources"] if p.name == "05_measurement.pdf")
        old_id = next(
            row["id"] for row in workspace.sources() if row["filename"] == changed.name
        )
        ingest_paths([changed], workspace)
        new_id = next(
            row["id"]
            for row in workspace.sources()
            if row["filename"] == changed.name and row["id"] != old_id
        )
        edited_ids = [new_id if sid == old_id else sid for sid in source_ids]
        cases[f"{workflow}/one-page-edit"] = measure(
            workspace, provider, edited_ids, workflow
        )
        if workflow == "assigned":
            cases["assigned/continuity-review"] = measure(
                workspace,
                provider,
                edited_ids,
                workflow,
                extra_options={"document_review": True},
            )
            for fmt in ("assessment", "podcast-script"):
                cases[f"assigned/{fmt}"] = measure(
                    workspace,
                    provider,
                    edited_ids,
                    workflow,
                    extra_options={"format": fmt},
                )
    task_workspace = Workspace(folder / "task-reading")
    ingest_paths(corpus["sources"], task_workspace)
    task_source_ids = [row["id"] for row in task_workspace.sources()]
    task_provider = ScriptedProvider()
    cases["planned/task-reading-cold"] = measure(
        task_workspace,
        task_provider,
        task_source_ids,
        "planned",
        extra_options={"reading": "task"},
    )
    cases["planned/task-reading-new-brief"] = measure(
        task_workspace,
        task_provider,
        task_source_ids,
        "planned",
        brief="Create a concise reference preserving every fixture fact and qualification.",
        extra_options={"reading": "task"},
    )
    for fault in ("one-invalid-quote", "persistent-invalid-quote"):
        workspace = Workspace(folder / fault)
        ingest_paths(corpus["sources"], workspace)
        source_ids = [row["id"] for row in workspace.sources()]
        cases[f"planned/{fault}"] = measure(
            workspace, ScriptedProvider(fault=fault), source_ids, "planned"
        )
    unicode_facts = [
        "[U01] 风机在测试模式下继续运行；测试结束后停机。",
        "[U02] परीक्षण पूरा होने पर रोटर रुकता है।",
        "[U03] يتوقف الدوار بعد انتهاء الاختبار.",
        "[U04] Keep mM, MM, −2, -2, 2–3 and 2-3 distinct in the original record.",
    ]
    unicode_source = folder / "dense-unicode.md"
    unicode_source.write_text(
        "# Dense Unicode fixture\n\n"
        + unicode_facts[0] * 100
        + "\n\n"
        + "\n\n".join(unicode_facts[1:]),
        encoding="utf-8",
    )
    for workflow in ("direct", "planned"):
        workspace = Workspace(folder / f"unicode-{workflow}")
        ingest_paths([unicode_source], workspace)
        cases[f"{workflow}/dense-unicode"] = measure(
            workspace,
            ScriptedProvider(),
            [s["id"] for s in workspace.sources()],
            workflow,
            required_facts=unicode_facts,
        )
    wide_facts = []
    wide_source = folder / "many-apparatus.md"
    paragraphs = ["# Original apparatus stress fixture"]
    for index in range(80):
        fact = f"[W{index:03d}] Apparatus M{index:03d} has its own inspection log and must retain that log when moved."
        wide_facts.append(fact)
        paragraphs.extend(
            [
                f"## [mechanics] [M{index:03d}] Separate apparatus record",
                fact
                + " This apparatus is identified separately from every other record in this fixture. Its test conditions and service decisions belong to its own record. Similar wording in another record does not authorize merging their identities. Preserve the apparatus identifier, the log requirement, and the movement condition together in the finished reference.",
            ]
        )
    wide_source.write_text("\n\n".join(paragraphs), encoding="utf-8")
    workspace = Workspace(folder / "wide-planned")
    ingest_paths([wide_source], workspace)
    cases["planned/bounded-hierarchy-and-writing"] = measure(
        workspace,
        ScriptedProvider(),
        [s["id"] for s in workspace.sources()],
        "planned",
        required_facts=wide_facts,
        extra_options={"max_input_bytes": 30_000},
    )
    return {
        "schema": "lamina-workflow-matrix-1",
        "evidence_kind": "deterministic fixture, synthetic variable delays, actual parsers and engine",
        "corpus": parse_metrics,
        "latency_model": {
            "distribution": "lognormal",
            "log_sigma": 0.7,
            "median_write_ms": 8,
            "median_other_ms": 4,
            "seed": "SHA256 of each exact request",
            "workers": 4,
        },
        "limits": [
            "The scripted provider echoes source text and knows the qualification oracle.",
            "No model quality, live latency, monetary cost, or universal optimum is measured.",
            "First review completion can precede a required repair; it is not time to accepted section.",
            "Citations do not establish entailment. Unit citation counts do not establish extraction recall.",
        ],
        "cases": cases,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--workdir", type=Path)
    args = parser.parse_args()

    def implementation_fingerprint():
        root = Path(__file__).resolve().parents[1]
        files = sorted((root / "src" / "lamina").rglob("*.py")) + [
            Path(__file__).resolve()
        ]
        checksum = hashlib.sha256()
        for path in files:
            checksum.update(str(path.relative_to(root)).encode())
            checksum.update(path.read_bytes())
        return checksum.hexdigest()

    before = implementation_fingerprint()
    if args.workdir:
        args.workdir.mkdir(parents=True, exist_ok=True)
        result = run_matrix(args.workdir)
    else:
        with tempfile.TemporaryDirectory(prefix="lamina-matrix-") as temporary:
            result = run_matrix(Path(temporary))
    result["implementation_sha256"] = implementation_fingerprint()
    result["implementation_unchanged_during_run"] = (
        before == result["implementation_sha256"]
    )
    result["fixture_assertions"] = {
        "unexpected_failures": [
            name
            for name, row in result["cases"].items()
            if row["status"] == "failed" and name != "planned/persistent-invalid-quote"
        ],
        "successful_outputs_missing_required_facts": [
            name
            for name, row in result["cases"].items()
            if row.get("fixture_oracle", {}).get("missing_fact_ids")
        ],
        "held_out_leaks": [
            name
            for name, row in result["cases"].items()
            if row.get("fixture_oracle", {}).get("held_out_key_present")
        ],
        "persistent_invalid_read_honestly_failed": result["cases"][
            "planned/persistent-invalid-quote"
        ]["status"]
        == "failed",
        "warm_runs_use_no_provider_calls": all(
            result["cases"][f"{workflow}/warm"]["calls"] == 0
            for workflow in ("direct", "planned", "assigned")
        ),
        "capacity_respected": all(
            row["peak_simultaneous_provider_calls"] <= 4
            for row in result["cases"].values()
        ),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                key: {
                    "status": row["status"],
                    "calls": row["calls"],
                    "error": row["error"],
                }
                for key, row in result["cases"].items()
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
