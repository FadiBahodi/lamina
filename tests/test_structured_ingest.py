import shutil
import subprocess

import pytest

from lamina.ingest import read_source


def test_markdown_fences_tables_and_reference_definitions_survive(tmp_path):
    source = tmp_path / "structured.md"
    table = "| Design | Condition |\n| --- | --- |\n" + "| Blade | Low wind |\n" * 800
    source.write_text(
        "# Wind\n\n```python\n# This is code\nprint('wind')\n```\n\n"
        + table
        + "\n[ref]: https://example.org\n"
    )
    _, units = read_source(source)
    assert all(u["heading"] == "Wind" for u in units)
    assert any(table.strip() in u["text"] for u in units)
    assert "[ref]: https://example.org" in "\n".join(u["text"] for u in units)
    assert "# This is code" in units[0]["text"]


def test_pptx_keeps_table_rows_and_speaker_notes(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Blade designs"
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(5), Inches(2)
    ).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Design", "Condition"
    table.cell(1, 0).text, table.cell(1, 1).text = "Wide blade", "Low wind"
    slide.notes_slide.notes_text_frame.text = "Do not generalize to every turbine."
    path = tmp_path / "wind.pptx"
    deck.save(path)
    _, units = read_source(path)
    assert any("Design | Condition\nWide blade | Low wind" == u["text"] for u in units)
    assert any(
        "speaker notes" in u["locator"] and "Do not generalize" in u["text"]
        for u in units
    )
    assert {u["structural_group"] for u in units} == {"slide:1"}
    assert {u["kind"] for u in units} == {"text", "table", "speaker_notes"}


def test_markdown_retains_heading_ancestry_and_whole_blocks(tmp_path):
    source = tmp_path / "wind.md"
    paragraph = "风使叶片旋转。" * 4000
    source.write_text(
        "# Wind\n\n## Rotor\n\nA rotor transfers torque.\n\n"
        "### Limits\n\n" + paragraph + "\n\n- First condition.\n- Second condition.\n"
    )
    _, units = read_source(source)
    assert [u["kind"] for u in units] == ["paragraph", "paragraph", "bullet_list"]
    assert units[1]["text"] == paragraph
    assert units[1]["heading_path"] == ["Wind", "Rotor", "Limits"]
    assert units[1]["section_id"] == units[2]["section_id"]
    assert units[0]["section_id"] != units[1]["section_id"]


def test_markdown_reference_definition_resolves_across_sections(tmp_path):
    source = tmp_path / "wind.md"
    source.write_text(
        "# Wind\n\n## Design\n\nSee [blade dimensions][spec].\n\n"
        '## References\n\n[spec]: https://example.org/blades "Blade dimensions"\n'
    )
    _, units = read_source(source)
    assert units[0]["reference_context"] == [
        {
            "label": "SPEC",
            "href": "https://example.org/blades",
            "title": "Blade dimensions",
        }
    ]
    assert any("[spec]: https://example.org/blades" in u["text"] for u in units)


def test_each_slide_has_an_explicit_group_and_order(tmp_path):
    pptx = pytest.importorskip("pptx")
    deck = pptx.Presentation()
    for title, note in [
        ("Wind load", "A rotating blade sees changing load."),
        ("Continuation", "The preceding load changes the shaft torque."),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = note
    path = tmp_path / "sequence.pptx"
    deck.save(path)
    _, units = read_source(path)
    assert [u["slide"] for u in units] == [1, 1, 2, 2]
    assert units[0]["structural_group"] == units[1]["structural_group"]
    assert units[1]["structural_group"] != units[2]["structural_group"]
    assert "preceding load" in units[3]["text"]


def test_local_ocr_preserves_original_identity(tmp_path):
    if not shutil.which("ocrmypdf"):
        pytest.skip("OCRmyPDF is optional")
    pytest.importorskip("pypdf")
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1800, 600), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("DejaVuSans.ttf", 54)
    draw.text((100, 200), "Wind turns the blades.", fill="black", font=font)
    source = tmp_path / "scan.pdf"
    image.save(source, resolution=200)
    original = source.read_bytes()
    with pytest.raises(ValueError, match="no extractable text"):
        read_source(source)
    info, units = read_source(source, ocr=True)
    assert "Wind turns the blades" in " ".join(u["text"] for u in units)
    assert info["ocr"] is True
    assert source.read_bytes() == original
